from .browser_runtime import launch as launch_browser
"""Local structured office readers and authorized cloud export fallback."""
from pathlib import Path
import json
from .wps_reader import DocumentReadError


def spreadsheet_blocks(path):
    from openpyxl import load_workbook
    book=load_workbook(path,data_only=False);values=load_workbook(path,data_only=True)
    result=[];images=[]
    for sheet in book:
        if sheet.max_row*sheet.max_column>2_000_000:raise DocumentReadError('工作表范围过大，尚未完成全表读取')
        rows=[]
        for row in sheet.iter_rows():
            cells=[]
            for cell in row:
                if cell.value is None:continue
                text=str(cell.value)
                if cell.data_type=='f':
                    cached=values[sheet.title][cell.coordinate].value
                    text+='；缓存结果：'+(str(cached) if cached is not None else '未提供，不推算结果')
                cells.append(cell.coordinate+'='+text)
            if cells:rows.append(' | '.join(cells))
        merges=', '.join(str(r) for r in sheet.merged_cells.ranges)
        prefix='工作表：'+sheet.title+'（'+sheet.sheet_state+'）\n'+('合并单元格：'+merges+'\n' if merges else '')
        # Repeat sheet identity for each bounded block; coordinates remain stable evidence.
        batch=[];size=0
        for row in rows:
            if size+len(row)>3500 and batch:result.append((len(result)+1,'sheet',prefix+'\n'.join(batch)));batch=[];size=0
            batch.append(row);size+=len(row)
        if batch:result.append((len(result)+1,'sheet',prefix+'\n'.join(batch)))
        for image in sheet._images:images.append((sheet.title,image._data()))
        for chart in sheet._charts:
            tree=chart.to_tree()
            texts=[e.text for e in tree.iter() if str(e.tag).split('}')[-1] in ('t','f','v') and e.text]
            result.append((len(result)+1,'sheet_chart',prefix+'图表类型：'+type(chart).__name__+'\n图表标题、系列引用及显式值（引用对应本工作表单元格）：\n'+'\n'.join(texts)))
    book.close();values.close()
    if not result:raise DocumentReadError('工作簿正文为空')
    return result,images


def cloud_export(url,secret_file,folder,kind):
    """Use normal visible download actions. Never forge an export permission."""
    from playwright.sync_api import sync_playwright
    from urllib.parse import urlsplit
    from .wps_reader import validate_url,session_value
    validate_url(url);folder=Path(folder);folder.mkdir(parents=True,exist_ok=True)
    with sync_playwright() as p:
        b=launch_browser(p)
        try:
            c=b.new_context(accept_downloads=True)
            c.add_cookies([{'name':'wps_sid','value':session_value(secret_file),'domain':'.kdocs.cn','path':'/','secure':True}])
            page=c.new_page();page.goto(url,wait_until='domcontentloaded',timeout=45000);page.wait_for_timeout(8000)
            def visible(label,top_only=False):
                matches=page.get_by_text(label,exact=True)
                for i in range(matches.count()):
                    loc=matches.nth(i)
                    if loc.is_visible() and (not top_only or loc.bounding_box()['y']<200):return loc
                return None
            menu=visible('文件',True) or visible('更多',True)
            if menu:menu.click(timeout=3000);page.wait_for_timeout(500)
            action=visible('下载原文件') or visible('下载原件') or visible('下载')
            if not action:raise DocumentReadError('当前云文档未提供可用的授权导出入口，等待对应格式读取器验证')
            with page.expect_download(timeout=20000) as pending:action.click(timeout=3000)
            download=pending.value
            suffix=Path(download.suggested_filename).suffix.lower()
            expected={'.xlsx'} if kind=='spreadsheet' else {'.md','.markdown','.txt'}
            if suffix not in expected:raise DocumentReadError('云文档导出格式与正文类型不符')
            target=folder/('cloud-source'+suffix);download.save_as(str(target))
            if target.stat().st_size>100_000_000:raise DocumentReadError('云文档导出超过本机读取上限')
            return target
        finally:b.close()
